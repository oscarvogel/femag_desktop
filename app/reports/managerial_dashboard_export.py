from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.config.database import database_proxy
from app.models.load_orders import LoadOrder
from app.models.masters import Carrier, Client, Driver, Product, Truck
from app.reports.managerial_dashboard import ManagerialDashboardService, ReportPeriod, _variation_percent

SHEET_README = "00_LEEME_IA"
SHEET_KPIS = "KPIs"
SHEET_EVOLUTION = "Evolucion_mensual"
SHEET_TOP_CLIENTS = "Top_clientes"
SHEET_TOP_PRODUCTS = "Top_productos"
SHEET_CLIENTS = "Clientes"
SHEET_PRODUCTS = "Articulos"
SHEET_LOADS = "Cargas"
SHEET_LINES = "Lineas"
SHEET_LOADS_PREV = "Cargas_mes_anterior"
SHEET_LINES_PREV = "Lineas_mes_anterior"
SHEET_REPURCHASE = "Recompra_cliente_articulo"
SHEET_CARRIERS = "Transportistas"

ALL_SHEETS = (
    SHEET_README,
    SHEET_KPIS,
    SHEET_EVOLUTION,
    SHEET_TOP_CLIENTS,
    SHEET_TOP_PRODUCTS,
    SHEET_CLIENTS,
    SHEET_PRODUCTS,
    SHEET_LOADS,
    SHEET_LINES,
    SHEET_LOADS_PREV,
    SHEET_LINES_PREV,
    SHEET_REPURCHASE,
    SHEET_CARRIERS,
)


def _order_map(period: ReportPeriod, effective_statuses: tuple[str, ...]) -> dict[int, dict]:
    if database_proxy.obj is None:
        return {}
    rows = (
        LoadOrder.select(
            LoadOrder.id,
            LoadOrder.order_number,
            LoadOrder.date,
            LoadOrder.status,
            LoadOrder.carrier,
            LoadOrder.driver,
            LoadOrder.truck,
            LoadOrder.client,
        )
        .where(
            LoadOrder.status.in_(effective_statuses),
            LoadOrder.date.between(period.start, period.end),
        )
        .dicts()
    )
    out: dict[int, dict] = {}
    for r in rows:
        out[int(r["id"])] = {
            "order_id": int(r["id"]),
            "order_number": r["order_number"],
            "date": r["date"],
            "status": r["status"],
            "carrier_id": r["carrier"],
            "driver_id": r["driver"],
            "truck_id": r["truck"],
            "header_client_id": r["client"],
        }
    return out


def _names_by_id(model, ids: set[int], field: str = "name") -> dict[int, str]:
    if not ids or database_proxy.obj is None:
        return {}
    return {r["id"]: r[field] for r in model.select(model.id, getattr(model, field)).where(model.id.in_(sorted(ids))).dicts()}


def _client_cuits(ids: set[int]) -> dict[int, str | None]:
    if not ids or database_proxy.obj is None:
        return {}
    return {r["id"]: r["cuit"] for r in Client.select(Client.id, Client.cuit).where(Client.id.in_(sorted(ids))).dicts()}


def _product_meta(ids: set[int]) -> dict[int, dict]:
    if not ids or database_proxy.obj is None:
        return {}
    out = {}
    for r in Product.select(Product.id, Product.codigo, Product.name, Product.unit).where(Product.id.in_(sorted(ids))).dicts():
        out[int(r["id"])] = {"codigo": r["codigo"], "name": r["name"], "unit": r["unit"]}
    return out


def _carrier_meta(ids: set[int]) -> dict[int, dict]:
    if not ids or database_proxy.obj is None:
        return {}
    out = {}
    for r in Carrier.select(Carrier.id, Carrier.name, Carrier.cuit).where(Carrier.id.in_(sorted(ids))).dicts():
        out[int(r["id"])] = {"name": r["name"], "cuit": r["cuit"]}
    return out


def build_export_datasets(
    service: ManagerialDashboardService, period: ReportPeriod
) -> dict[str, tuple[tuple[str, ...], list[tuple]]]:
    snapshot = service.snapshot(period)
    previous = snapshot.previous_period
    breakdown = service._dispatch_breakdown(period)
    prev_breakdown = service._dispatch_breakdown(previous)

    orders = _order_map(period, service.effective_statuses)
    prev_orders = _order_map(previous, service.effective_statuses)

    client_ids = {int(r["client_id"]) for r in breakdown} | {int(r["client_id"]) for r in prev_breakdown}
    product_ids = {int(r["product_id"]) for r in breakdown} | {int(r["product_id"]) for r in prev_breakdown}
    cuits = _client_cuits(client_ids)
    prod_meta = _product_meta(product_ids)

    carrier_ids = {v["carrier_id"] for v in orders.values() if v["carrier_id"]} | {
        v["carrier_id"] for v in prev_orders.values() if v["carrier_id"]
    }
    carrier_meta = _carrier_meta({int(c) for c in carrier_ids if c})
    driver_ids = {v["driver_id"] for v in orders.values() if v["driver_id"]} | {
        v["driver_id"] for v in prev_orders.values() if v["driver_id"]
    }
    truck_ids = {v["truck_id"] for v in orders.values() if v["truck_id"]} | {
        v["truck_id"] for v in prev_orders.values() if v["truck_id"]
    }
    driver_names = _names_by_id(Driver, driver_ids)
    truck_domains = _names_by_id(Truck, truck_ids, field="domain")

    def iso(d) -> str:
        return d.isoformat() if hasattr(d, "isoformat") else str(d or "")

    # ---- Clientes: todos con movimiento ----
    cur_c: dict[int, dict] = defaultdict(lambda: {"total": 0.0, "kilos": 0.0, "qty": 0.0, "orders": set(), "name": ""})
    for r in breakdown:
        cid = int(r["client_id"])
        cur_c[cid]["name"] = r["client_name"]
        cur_c[cid]["total"] += float(r["total"] or 0)
        cur_c[cid]["kilos"] += float(r["kilos"] or 0)
        cur_c[cid]["qty"] += float(r["quantity"] or 0)
        cur_c[cid]["orders"].add(int(r["order_id"]))
    prev_c: dict[int, dict] = defaultdict(lambda: {"total": 0.0, "kilos": 0.0, "orders": set()})
    for r in prev_breakdown:
        cid = int(r["client_id"])
        prev_c[cid]["total"] += float(r["total"] or 0)
        prev_c[cid]["kilos"] += float(r["kilos"] or 0)
        prev_c[cid]["orders"].add(int(r["order_id"]))
    total_all = sum(v["total"] for v in cur_c.values()) or 1.0
    clients_rows = []
    for cid in sorted(cur_c, key=lambda c: cur_c[c]["total"], reverse=True):
        c = cur_c[cid]
        p = prev_c.get(cid, {"total": 0.0, "kilos": 0.0, "orders": set()})
        clients_rows.append((
            cid, c["name"], cuits.get(cid) or "",
            len(c["orders"]), round(c["qty"], 3),
            round(c["total"], 2), round(c["kilos"] / 1000.0, 3),
            round(p["total"], 2), round(p["kilos"] / 1000.0, 3),
            _variation_percent(c["total"], p["total"]),
            round(c["total"] / total_all * 100, 2),
            round(c["total"] / len(c["orders"]), 2) if c["orders"] else 0.0,
        ))
    clients_headers = ("cliente_id", "cliente", "cuit", "cargas", "cantidad", "valorizado",
                       "toneladas", "valorizado_mes_ant", "toneladas_mes_ant",
                       "variacion_%", "participacion_%", "ticket_promedio")

    # ---- Artículos: todos con movimiento ----
    cur_p: dict[int, dict] = defaultdict(lambda: {"total": 0.0, "kilos": 0.0, "qty": 0.0, "orders": set(), "clients": set(), "name": ""})
    for r in breakdown:
        pid = int(r["product_id"])
        cur_p[pid]["name"] = r["product_name"]
        cur_p[pid]["total"] += float(r["total"] or 0)
        cur_p[pid]["kilos"] += float(r["kilos"] or 0)
        cur_p[pid]["qty"] += float(r["quantity"] or 0)
        cur_p[pid]["orders"].add(int(r["order_id"]))
        cur_p[pid]["clients"].add(int(r["client_id"]))
    prev_p: dict[int, dict] = defaultdict(lambda: {"total": 0.0, "kilos": 0.0})
    for r in prev_breakdown:
        pid = int(r["product_id"])
        prev_p[pid]["total"] += float(r["total"] or 0)
        prev_p[pid]["kilos"] += float(r["kilos"] or 0)
    products_rows = []
    for pid in sorted(cur_p, key=lambda c: cur_p[c]["total"], reverse=True):
        c = cur_p[pid]
        meta = prod_meta.get(pid, {})
        p = prev_p.get(pid, {"total": 0.0, "kilos": 0.0})
        products_rows.append((
            pid, meta.get("codigo") or "", c["name"], meta.get("unit") or "",
            len(c["orders"]), len(c["clients"]), round(c["qty"], 3),
            round(c["total"], 2), round(c["kilos"] / 1000.0, 3),
            round(p["total"], 2), _variation_percent(c["total"], p["total"]),
            round(c["total"] / total_all * 100, 2),
        ))
    products_headers = ("articulo_id", "codigo", "articulo", "unidad", "cargas", "clientes",
                        "cantidad", "valorizado", "toneladas", "valorizado_mes_ant",
                        "variacion_%", "participacion_%")

    # ---- Cargas: una fila por orden efectiva ----
    order_totals: dict[int, dict] = defaultdict(lambda: {"total": 0.0, "kilos": 0.0, "lines": 0, "clients": set()})
    for r in breakdown:
        oid = int(r["order_id"])
        order_totals[oid]["total"] += float(r["total"] or 0)
        order_totals[oid]["kilos"] += float(r["kilos"] or 0)
        order_totals[oid]["lines"] += 1
        order_totals[oid]["clients"].add(r["client_name"])
    loads_rows = []
    for oid, o in sorted(orders.items(), key=lambda kv: (str(kv[1]["date"]), kv[1]["order_number"])):
        t = order_totals.get(oid, {"total": 0.0, "kilos": 0.0, "lines": 0, "clients": set()})
        cm = carrier_meta.get(int(o["carrier_id"])) if o["carrier_id"] else {}
        loads_rows.append((
            o["order_number"], iso(o["date"]), o["status"],
            (cm or {}).get("name", ""), driver_names.get(o["driver_id"], ""),
            truck_domains.get(o["truck_id"], ""),
            "; ".join(sorted(t["clients"])), t["lines"],
            round(t["total"], 2), round(t["kilos"] / 1000.0, 3),
            round(t["total"] / t["lines"], 2) if t["lines"] else 0.0,
        ))
    loads_headers = ("carga_nro", "fecha", "estado", "transportista", "chofer", "camion",
                     "clientes", "lineas", "valorizado", "toneladas", "valorizado_por_linea")

    # ---- Líneas: tabla de hechos granular para IA ----
    lines_rows = []
    for r in sorted(breakdown, key=lambda x: (str(orders.get(int(x["order_id"]), {}).get("date", "")), int(x["order_id"]))):
        oid = int(r["order_id"])
        o = orders.get(oid, {})
        cm = carrier_meta.get(int(o["carrier_id"])) if o.get("carrier_id") else {}
        pid = int(r["product_id"])
        meta = prod_meta.get(pid, {})
        qty = float(r["quantity"] or 0)
        kilos = float(r["kilos"] or 0)
        total = float(r["total"] or 0)
        lines_rows.append((
            iso(o.get("date", "")), o.get("order_number", ""), o.get("status", ""),
            (cm or {}).get("name", ""), driver_names.get(o.get("driver_id", 0), ""),
            truck_domains.get(o.get("truck_id", 0), ""),
            int(r["client_id"]), r["client_name"], cuits.get(int(r["client_id"])) or "",
            pid, meta.get("codigo") or "", r["product_name"], meta.get("unit") or "",
            round(qty, 3), round(kilos / 1000.0, 3), round(total, 2),
            round(total / qty, 2) if qty else 0.0,
        ))
    lines_headers = ("fecha", "carga_nro", "estado", "transportista", "chofer", "camion",
                     "cliente_id", "cliente", "cliente_cuit",
                     "articulo_id", "codigo", "articulo", "unidad",
                     "cantidad", "toneladas", "valorizado", "precio_unitario")

    # ---- Cargas y líneas del período anterior (mismo corte) ----
    prev_order_totals: dict[int, dict] = defaultdict(lambda: {"total": 0.0, "kilos": 0.0, "lines": 0, "clients": set()})
    for r in prev_breakdown:
        oid = int(r["order_id"])
        prev_order_totals[oid]["total"] += float(r["total"] or 0)
        prev_order_totals[oid]["kilos"] += float(r["kilos"] or 0)
        prev_order_totals[oid]["lines"] += 1
        prev_order_totals[oid]["clients"].add(r["client_name"])
    prev_loads_rows = []
    for oid, o in sorted(prev_orders.items(), key=lambda kv: (str(kv[1]["date"]), kv[1]["order_number"])):
        t = prev_order_totals.get(oid, {"total": 0.0, "kilos": 0.0, "lines": 0, "clients": set()})
        cm = carrier_meta.get(int(o["carrier_id"])) if o["carrier_id"] else {}
        prev_loads_rows.append((
            o["order_number"], iso(o["date"]), o["status"],
            (cm or {}).get("name", ""), driver_names.get(o["driver_id"], ""),
            truck_domains.get(o["truck_id"], ""),
            "; ".join(sorted(t["clients"])), t["lines"],
            round(t["total"], 2), round(t["kilos"] / 1000.0, 3),
            round(t["total"] / t["lines"], 2) if t["lines"] else 0.0,
        ))

    prev_lines_rows = []
    for r in sorted(prev_breakdown, key=lambda x: (str(prev_orders.get(int(x["order_id"]), {}).get("date", "")), int(x["order_id"]))):
        oid = int(r["order_id"])
        o = prev_orders.get(oid, {})
        cm = carrier_meta.get(int(o["carrier_id"])) if o.get("carrier_id") else {}
        pid = int(r["product_id"])
        meta = prod_meta.get(pid, {})
        qty = float(r["quantity"] or 0)
        kilos = float(r["kilos"] or 0)
        total = float(r["total"] or 0)
        prev_lines_rows.append((
            iso(o.get("date", "")), o.get("order_number", ""), o.get("status", ""),
            (cm or {}).get("name", ""), driver_names.get(o.get("driver_id", 0), ""),
            truck_domains.get(o.get("truck_id", 0), ""),
            int(r["client_id"]), r["client_name"], cuits.get(int(r["client_id"])) or "",
            pid, meta.get("codigo") or "", r["product_name"], meta.get("unit") or "",
            round(qty, 3), round(kilos / 1000.0, 3), round(total, 2),
            round(total / qty, 2) if qty else 0.0,
        ))

    # ---- Recompra cliente x artículo (criterio principal: kilos) ----
    repurchase = service.product_client_repurchase(period, previous)
    repurchase_rows = []
    for r in repurchase:
        meta = prod_meta.get(int(r["product_id"]), {})
        repurchase_rows.append((
            int(r["client_id"]), r["client_name"], cuits.get(int(r["client_id"])) or "",
            int(r["product_id"]), meta.get("codigo") or "", r["product_name"],
            r["previous_quantity"], r["previous_tonnes"], r["previous_total"],
            r["current_quantity"], r["current_tonnes"], r["current_total"],
            r["diff_tonnes"], r["variation_kilos_percent"], r["variation_total_percent"],
            r["status"],
        ))
    repurchase_headers = ("cliente_id", "cliente", "cuit",
                          "articulo_id", "codigo", "articulo",
                          "cantidad_mes_ant", "toneladas_mes_ant", "valorizado_mes_ant",
                          "cantidad_actual", "toneladas_actual", "valorizado_actual",
                          "dif_toneladas", "variacion_kilos_%", "variacion_valorizado_%",
                          "clasificacion")

    # ---- Transportistas ----
    cur_t: dict[int, dict] = defaultdict(lambda: {"total": 0.0, "kilos": 0.0, "orders": set(), "clients": set()})
    for r in breakdown:
        o = orders.get(int(r["order_id"]))
        if not o or not o.get("carrier_id"):
            continue
        tid = int(o["carrier_id"])
        cur_t[tid]["total"] += float(r["total"] or 0)
        cur_t[tid]["kilos"] += float(r["kilos"] or 0)
        cur_t[tid]["orders"].add(int(r["order_id"]))
        cur_t[tid]["clients"].add(int(r["client_id"]))
    prev_t: dict[int, dict] = defaultdict(lambda: {"total": 0.0, "kilos": 0.0})
    for r in prev_breakdown:
        o = prev_orders.get(int(r["order_id"]))
        if not o or not o.get("carrier_id"):
            continue
        tid = int(o["carrier_id"])
        prev_t[tid]["total"] += float(r["total"] or 0)
        prev_t[tid]["kilos"] += float(r["kilos"] or 0)
    carriers_rows = []
    for tid in sorted(cur_t, key=lambda t: cur_t[t]["total"], reverse=True):
        c = cur_t[tid]
        meta = carrier_meta.get(tid, {})
        p = prev_t.get(tid, {"total": 0.0, "kilos": 0.0})
        carriers_rows.append((
            tid, meta.get("name", ""), meta.get("cuit") or "",
            len(c["orders"]), len(c["clients"]),
            round(c["total"], 2), round(c["kilos"] / 1000.0, 3),
            round(p["total"], 2), _variation_percent(c["total"], p["total"]),
            round(c["total"] / total_all * 100, 2),
            round(c["total"] / len(c["orders"]), 2) if c["orders"] else 0.0,
        ))
    carriers_headers = ("transportista_id", "transportista", "cuit", "cargas", "clientes",
                        "valorizado", "toneladas", "valorizado_mes_ant",
                        "variacion_%", "participacion_%", "valorizado_por_carga")

    # ---- KPIs / evolución / tops (reusan snapshot) ----
    kpi_rows = [
        ("despachos_valorizados", snapshot.valued_dispatches.current, snapshot.valued_dispatches.previous,
         snapshot.valued_dispatches.variation_percent),
        ("toneladas", snapshot.tonnes.current, snapshot.tonnes.previous, snapshot.tonnes.variation_percent),
        ("ordenes_cargas", snapshot.orders.current, snapshot.orders.previous, snapshot.orders.variation_percent),
        ("ticket_promedio", snapshot.average_ticket.current, snapshot.average_ticket.previous,
         snapshot.average_ticket.variation_percent),
        ("saldo_clientes", snapshot.total_receivables, "", ""),
        ("saldo_vencido", snapshot.overdue_receivables, "", ""),
    ]
    kpi_headers = ("metrica", "actual", "mes_anterior_mismo_corte", "variacion_%")

    evo_rows = [(
        r["label"], r["total"], r.get("previous_total", 0.0), r.get("variation_percent"),
        r["tonnes"], r.get("previous_tonnes", 0.0), r.get("variation_tonnes"),
        r["orders"], r.get("previous_orders", 0), r.get("previous_range", ""),
    ) for r in snapshot.monthly_evolution]
    evo_headers = ("mes", "valorizado", "valorizado_mes_ant", "variacion_%",
                   "toneladas", "toneladas_mes_ant", "variacion_tn_%",
                   "ordenes", "ordenes_mes_ant", "rango_comparable")

    top_c_rows = [(i + 1, r["name"], r["total"], r.get("previous_total", 0.0),
                   r.get("variation_percent"), r["tonnes"]) for i, r in enumerate(snapshot.top_clients)]
    top_c_headers = ("ranking", "cliente", "valorizado", "valorizado_mes_ant", "variacion_%", "toneladas")

    top_p_rows = [(i + 1, r["name"], r["total"], r.get("previous_total", 0.0),
                   r.get("variation_percent"), r["tonnes"]) for i, r in enumerate(snapshot.top_products)]
    top_p_headers = ("ranking", "articulo", "valorizado", "valorizado_mes_ant", "variacion_%", "toneladas")

    readme_rows = [
        ("titulo", "Dashboard gerencial FEMAG - export para analisis IA"),
        ("generado_en", datetime.now().isoformat(timespec="seconds")),
        ("periodo_actual", snapshot.period.describe_range()),
        ("periodo_anterior_equivalente", previous.describe_range()),
        ("etiqueta_periodo", snapshot.period.label),
        ("criterio", "Despachos efectivos: " + ", ".join(snapshot.effective_statuses) + ". Comparacion mismo corte dia por dia."),
        ("moneda", snapshot.currency),
        ("hoja_KPIs", "Metricas del periodo vs anterior mismo corte."),
        ("hoja_Evolucion_mensual", "12 meses; mes parcial vs mismo corte del mes anterior."),
        ("hoja_Top_clientes", "Top 10 por valorizado con comparativa."),
        ("hoja_Top_productos", "Top 10 por kilos/valorizado con comparativa."),
        ("hoja_Clientes", "TODOS los clientes con movimiento en el periodo, ordenados por valorizado."),
        ("hoja_Articulos", "TODOS los articulos con movimiento en el periodo."),
        ("hoja_Cargas", "Detalle por carga del periodo actual: una fila por orden efectiva."),
        ("hoja_Lineas", "Tabla de hechos granular actual: una fila por carga x cliente x articulo. Ideal para IA."),
        ("hoja_Cargas_mes_anterior", "Detalle por carga del periodo anterior equivalente (mismo corte)."),
        ("hoja_Lineas_mes_anterior", "Detalle granular del periodo anterior: permite ver que compro cada cliente por articulo."),
        ("hoja_Recompra_cliente_articulo", "Comparativa cliente x articulo por KILOS: sin_recompra, compro_menos, igual, compro_mas, nuevo."),
        ("hoja_Transportistas", "Detalle por transportista agregado, con comparativa mes anterior mismo corte."),
        ("fechas", "Formato ISO YYYY-MM-DD para analisis automatico."),
        ("numeros", "Montos en ARS como numero, toneladas con 3 decimales, variacion en %."),
    ]

    return {
        SHEET_README: (("campo", "valor"), readme_rows),
        SHEET_KPIS: (kpi_headers, kpi_rows),
        SHEET_EVOLUTION: (evo_headers, evo_rows),
        SHEET_TOP_CLIENTS: (top_c_headers, top_c_rows),
        SHEET_TOP_PRODUCTS: (top_p_headers, top_p_rows),
        SHEET_CLIENTS: (clients_headers, clients_rows),
        SHEET_PRODUCTS: (products_headers, products_rows),
        SHEET_LOADS: (loads_headers, loads_rows),
        SHEET_LINES: (lines_headers, lines_rows),
        SHEET_LOADS_PREV: (loads_headers, prev_loads_rows),
        SHEET_LINES_PREV: (lines_headers, prev_lines_rows),
        SHEET_REPURCHASE: (repurchase_headers, repurchase_rows),
        SHEET_CARRIERS: (carriers_headers, carriers_rows),
    }


def export_dashboard_xlsx(
    service: ManagerialDashboardService, period: ReportPeriod, path: str | Path
) -> Path:
    path = Path(path)
    datasets = build_export_datasets(service, period)
    wb = Workbook()
    first = True
    for sheet_name in ALL_SHEETS:
        headers, rows = datasets[sheet_name]
        ws = wb.active if first else wb.create_sheet(sheet_name)
        ws.title = sheet_name
        first = False
        ws.append(list(headers))
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append(list(row))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
        for col_idx, header in enumerate(headers, start=1):
            letter = get_column_letter(col_idx)
            max_len = len(str(header))
            for cell in ws[letter][1:]:
                if cell.value is None:
                    continue
                max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[letter].width = min(max_len + 2, 42)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
