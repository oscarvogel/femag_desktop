from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from app.reports.product_profitability import ProfitabilitySnapshot


SUMMARY_HEADERS = ("Producto","Cantidad","Unidad","Venta neta","Venta con costo conocido","Costo conocido","Utilidad bruta","Margen %","Cobertura %")
DETAIL_HEADERS = ("Fecha","Orden","Cliente","Producto","Cantidad","Unidad","Precio venta neto","Costo aplicado","Venta neta","Costo total","Utilidad bruta","Margen %","Costo informado")


def export_profitability_xlsx(snapshot: ProfitabilitySnapshot, path: str | Path, *, start, end) -> Path:
    path=Path(path)
    wb=Workbook()
    ws=wb.active; ws.title="Resumen por producto"
    ws.append(["Rentabilidad por producto", f"{start:%d/%m/%Y} - {end:%d/%m/%Y}"])
    ws.append(["Ventas netas",snapshot.sales,"Costo conocido",snapshot.cost,"Utilidad bruta",snapshot.gross_profit,"Margen %",snapshot.margin_percent,"Cobertura %",snapshot.cost_coverage_percent])
    ws.append([])
    ws.append(SUMMARY_HEADERS)
    for cell in ws[4]: cell.font=Font(bold=True)
    for row in snapshot.products:
        ws.append([row.product,row.quantity,row.unit,row.sales,row.known_cost_sales,row.cost,row.gross_profit,row.margin_percent,row.cost_coverage_percent])
    detail=wb.create_sheet("Detalle por operacion")
    detail.append(DETAIL_HEADERS)
    for cell in detail[1]: cell.font=Font(bold=True)
    for row in snapshot.lines:
        detail.append([row.date,row.order_number,row.client,row.product,row.quantity,row.unit,row.sale_unit_price,row.applied_unit_cost,row.sale_amount,row.cost_amount,row.gross_profit,row.margin_percent,"Sí" if row.applied_unit_cost is not None else "No"])
    for sheet in wb.worksheets:
        sheet.freeze_panes="A5" if sheet is ws else "A2"
        for column in sheet.columns:
            width=min(max(len(str(cell.value or "")) for cell in column)+2,35)
            sheet.column_dimensions[column[0].column_letter].width=width
    wb.save(path)
    return path


def export_profitability_csv(snapshot: ProfitabilitySnapshot, path: str | Path) -> Path:
    path=Path(path)
    with path.open("w",newline="",encoding="utf-8-sig") as handle:
        writer=csv.writer(handle,delimiter=";")
        writer.writerow(DETAIL_HEADERS)
        for row in snapshot.lines:
            writer.writerow([row.date.strftime("%Y-%m-%d"),row.order_number,row.client,row.product,row.quantity,row.unit,row.sale_unit_price,row.applied_unit_cost if row.applied_unit_cost is not None else "",row.sale_amount,row.cost_amount if row.cost_amount is not None else "",row.gross_profit if row.gross_profit is not None else "",row.margin_percent if row.margin_percent is not None else "","Sí" if row.applied_unit_cost is not None else "No"])
    return path
