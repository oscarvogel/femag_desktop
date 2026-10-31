from __future__ import annotations

import os
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.models.masters import Client, Product, Salesperson
from app.config.database import database_proxy
from app.reports.managerial_dashboard import ManagerialDashboardService, ReportPeriod


def _client_cuits(ids: set[int]) -> dict[int, str | None]:
    if not ids or database_proxy.obj is None:
        return {}
    return {r["id"]: r["cuit"] for r in Client.select(Client.id, Client.cuit).where(Client.id.in_(sorted(ids))).dicts()}


def _product_meta(ids: set[int]) -> dict[int, dict]:
    if not ids or database_proxy.obj is None:
        return {}
    out = {}
    for r in Product.select(Product.id, Product.codigo, Product.name, Product.unit).where(Product.id.in_(sorted(ids))).dicts():
        out[int(r["id"])] = {"codigo": r["codigo"], "name": r["name"], "unit": r["unit"]}
    return out


def build_repurchase_datasets(
    service: ManagerialDashboardService,
    current: ReportPeriod,
    previous: ReportPeriod | None = None,
    *,
    product_id: int | None = None,
    salesperson_id: int | None = None,
) -> dict[str, tuple[tuple[str, ...], list[tuple]]]:
    previous = previous or current.previous_equivalent()
    rows = service.product_client_repurchase(
        current, previous, product_id=product_id, salesperson_id=salesperson_id
    )
    summary = service.product_repurchase_summary(
        current, previous, product_id=product_id, salesperson_id=salesperson_id
    )
    product_ids = {int(r["product_id"]) for r in rows}
    client_ids = {int(r["client_id"]) for r in rows}
    meta = _product_meta(product_ids)
    cuits = _client_cuits(client_ids)
    product_label = ""
    if product_id is not None and product_id in meta:
        code = meta[product_id].get("codigo") or ""
        product_label = f"{code} · {meta[product_id]['name']}" if code else meta[product_id]["name"]
    salesperson_label = "Todos los vendedores"
    if salesperson_id is not None and database_proxy.obj is not None:
        salesperson = Salesperson.get_or_none(Salesperson.id == salesperson_id)
        if salesperson is not None:
            salesperson_label = salesperson.name

    summary_rows = [
        ("producto", product_label or ("Todos los artículos" if product_id is None else str(product_id))),
        ("vendedor", salesperson_label),
        ("periodo_anterior", previous.describe_range()),
        ("periodo_actual", current.describe_range()),
        ("criterio", "Órdenes efectivas Emitida/Cerrada. Kilos físicos pallet+suelto con fallback. Sin duplicar por destino. Vendedor = asignación actual del cliente (proxy)."),
        ("tn_anterior", summary["previous_tonnes"]),
        ("tn_actual", summary["current_tonnes"]),
        ("variacion_tn", summary["diff_tonnes"]),
        ("variacion_pct", summary["variation_percent"]),
        ("valorizado_anterior", summary["previous_total"]),
        ("valorizado_actual", summary["current_total"]),
        ("clientes_compradores_anterior", summary["buyers_previous"]),
        ("clientes_compradores_actual", summary["buyers_current"]),
        ("clientes_sin_recompra", summary["count_sin_recompra"]),
        ("tn_sin_recompra", summary["lost_tonnes"]),
        ("clientes_redujeron", summary["count_redujeron"]),
        ("tn_reducciones", summary["reduction_tonnes"]),
        ("clientes_aumentaron", summary["count_aumentaron"]),
        ("tn_aumentos", summary["increase_tonnes"]),
        ("clientes_nuevos_periodo", summary["count_nuevos"]),
        ("tn_nuevos", summary["new_tonnes"]),
        ("conciliacion", f"{summary['reconciled_diff_tonnes']} TN = -{summary['lost_tonnes']} + ({summary['reduction_tonnes']}) + {summary['increase_tonnes']} + {summary['new_tonnes']}"),
        ("conciliacion_ok", "SI" if summary["reconciliation_ok"] else "NO"),
    ]

    def row_tuple(r: dict) -> tuple:
        m = meta.get(int(r["product_id"]), {})
        return (
            int(r["client_id"]), r["client_name"], cuits.get(int(r["client_id"])) or "",
            int(r["product_id"]), m.get("codigo") or "", r["product_name"],
            r.get("salesperson_name") or "",
            r["previous_last_date"] or "", r["current_last_date"] or "",
            r["previous_quantity"], r["previous_tonnes"], r["previous_total"], r["previous_orders"],
            r["current_quantity"], r["current_tonnes"], r["current_total"], r["current_orders"],
            r["diff_tonnes"], r["variation_kilos_percent"],
            r["status"],
        )

    headers = (
        "cliente_id", "cliente", "cuit",
        "articulo_id", "codigo", "articulo", "vendedor_asignacion_actual",
        "ultima_compra_anterior", "ultima_compra_actual",
        "cantidad_anterior", "tn_anterior", "valorizado_anterior", "ordenes_anterior",
        "cantidad_actual", "tn_actual", "valorizado_actual", "ordenes_actual",
        "dif_tn", "variacion_kilos_pct", "clasificacion",
    )
    by_status: dict[str, list[dict]] = {}
    for r in rows:
        by_status.setdefault(r["status"], []).append(r)

    # Detalle verificable: reutiliza breakdown (una fila por orden x destino x producto).
    verifiable: list[tuple] = []
    for period, tag in ((previous, "anterior"), (current, "actual")):
        for line in service._dispatch_breakdown(period):
            if product_id is not None and int(line["product_id"]) != product_id:
                continue
            if salesperson_id is not None and int(line.get("salesperson_id") or 0) != salesperson_id:
                continue
            pid = int(line["product_id"])
            m = meta.get(pid, {})
            verifiable.append((
                tag, str(line.get("order_date") or ""),
                int(line["order_id"]), int(line["client_id"]), line["client_name"],
                pid, m.get("codigo") or "", line["product_name"],
                round(float(line["quantity"] or 0), 3), round(float(line["kilos"] or 0) / 1000.0, 3),
                round(float(line["total"] or 0), 2),
            ))
    verifiable.sort(key=lambda t: (t[0], t[1], t[2]))
    detail_headers = (
        "periodo", "fecha_orden", "orden_id", "cliente_id", "cliente",
        "articulo_id", "codigo", "articulo", "cantidad", "toneladas", "valorizado",
    )
    return {
        "Resumen": (("campo", "valor"), summary_rows),
        "Sin_recompra": (headers, [row_tuple(r) for r in sorted(by_status.get("sin_recompra", []), key=lambda x: -x["previous_kilos"])]),
        "Reduccion": (headers, [row_tuple(r) for r in sorted(by_status.get("compro_menos", []), key=lambda x: x["diff_kilos"])]),
        "Aumento": (headers, [row_tuple(r) for r in sorted(by_status.get("compro_mas", []), key=lambda x: -x["diff_kilos"])]),
        "Incorporados": (headers, [row_tuple(r) for r in sorted(by_status.get("nuevo", []), key=lambda x: -x["current_kilos"])]),
        "Detalle": (detail_headers, verifiable),
    }


REPURCHASE_SHEETS = ("Resumen", "Sin_recompra", "Reduccion", "Aumento", "Incorporados", "Detalle")


def export_repurchase_xlsx(
    service: ManagerialDashboardService,
    current: ReportPeriod,
    path: str | Path,
    *,
    previous: ReportPeriod | None = None,
    product_id: int | None = None,
) -> Path:
    path = Path(path)
    datasets = build_repurchase_datasets(service, current, previous, product_id=product_id)
    wb = Workbook()
    first = True
    for sheet_name in REPURCHASE_SHEETS:
        headers, rows = datasets[sheet_name]
        ws = wb.active if first else wb.create_sheet(sheet_name)
        ws.title = sheet_name
        first = False
        ws.append(list(headers))
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append(list(row))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
        for col_idx, header in enumerate(headers, start=1):
            letter = get_column_letter(col_idx)
            max_len = len(str(header))
            for cell in ws[letter][1:]:
                if cell.value is None:
                    continue
                max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[letter].width = min(max_len + 2, 42)
def export_repurchase_xlsx(
    service: ManagerialDashboardService,
    current: ReportPeriod,
    path: str | Path,
    *,
    previous: ReportPeriod | None = None,
    product_id: int | None = None,
    salesperson_id: int | None = None,
    chart_image: str | Path | None = None,
) -> Path:
    path = Path(path)
    datasets = build_repurchase_datasets(
        service, current, previous, product_id=product_id, salesperson_id=salesperson_id
    )
    wb = Workbook()
    first = True
    for sheet_name in REPURCHASE_SHEETS:
        headers, rows = datasets[sheet_name]
        ws = wb.active if first else wb.create_sheet(sheet_name)
        ws.title = sheet_name
        first = False
        ws.append(list(headers))
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append(list(row))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
        for col_idx, header in enumerate(headers, start=1):
            letter = get_column_letter(col_idx)
            max_len = len(str(header))
            for cell in ws[letter][1:]:
                if cell.value is None:
                    continue
                max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[letter].width = min(max_len + 2, 42)
    if chart_image:
        try:
            from openpyxl.drawing.image import Image as XLImage

            png = Path(chart_image)
            if png.exists():
                ws = wb["Resumen"]
                start_row = ws.max_row + 3
                ws.cell(row=start_row, column=1, value="grafico_tn").font = Font(bold=True)
                img = XLImage(str(png))
                img.width = 640
                img.height = 360
                ws.add_image(img, f"B{start_row + 1}")
        except Exception:
            pass
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _svg_bars(prev_tn: float, cur_tn: float, prev_label: str, cur_label: str) -> str:
    maximum = max(prev_tn, cur_tn, 1.0)
    for step in (10, 20, 40, 80, 160, 320, 640):
        if maximum <= step:
            maximum = float(step)
            break
    else:
        maximum = maximum * 1.1
    width, height = 560, 260
    left, bottom, top = 56, 34, 14
    plot_h = height - top - bottom
    plot_w = width - left - 16
    grid = ""
    for frac in (0.0, 0.25, 0.5, 0.75, 1.0):
        y = top + plot_h * (1 - frac)
        grid += (
            f'<line x1="{left}" y1="{y:.0f}" x2="{width - 16}" y2="{y:.0f}" stroke="#e5e7eb"/>'
            f'<text x="{left - 8}" y="{y + 4:.0f}" text-anchor="end" font-size="11" fill="#6b7280">'
            f"{maximum * frac:.0f} tn</text>"
        )
    bars = ""
    slot = plot_w / 2.0
    bar_w = min(slot * 0.52, 120.0)
    down = cur_tn < prev_tn
    for index, value in enumerate((prev_tn, cur_tn)):
        frac = (value / maximum) if maximum else 0.0
        bar_h = plot_h * frac
        x = left + slot * index + (slot - bar_w) / 2.0
        y = top + plot_h - bar_h
        color = "#2e5fa3" if index == 0 else ("#c43b47" if down else "#11875d")
        label = prev_label if index == 0 else cur_label
        bars += (
            f'<rect x="{x:.0f}" y="{y:.0f}" width="{bar_w:.0f}" height="{max(bar_h, 2.0):.0f}" '
            f'rx="7" fill="{color}"/>'
            f'<text x="{x + bar_w / 2:.0f}" y="{max(y - 8, 12):.0f}" text-anchor="middle" '
            f'font-size="13" font-weight="700" fill="#111827">{value:,.2f}</text>'
            f'<text x="{x + bar_w / 2:.0f}" y="{height - 12}" text-anchor="middle" '
            f'font-size="12" fill="#374151">{label}</text>'
        )
    return (
        f'<svg viewBox="0 0 {width} {height}" width="100%" role="img" '
        f'aria-label="Comparativa de toneladas">'
        f'<rect x="0" y="0" width="{width}" height="{height}" fill="#ffffff"/>{grid}{bars}</svg>'
    )


def export_repurchase_html(
    service: ManagerialDashboardService,
    current: ReportPeriod,
    path: str | Path,
    *,
    previous: ReportPeriod | None = None,
    product_id: int | None = None,
    salesperson_id: int | None = None,
    product_label: str = "",
    salesperson_label: str = "",
) -> Path:
    previous = previous or current.previous_equivalent()
    doc = _render_repurchase_html(
        service,
        current,
        previous,
        product_id,
        product_label,
        salesperson_id=salesperson_id,
        salesperson_label=salesperson_label,
    )
    out = Path(path)
    if out.suffix.lower() != ".html":
        out = out.with_suffix(".html")
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(doc, encoding="utf-8")
    return out


class ManagerialRepurchaseHtmlReport:
    """Informe web de recompra, con el mismo flujo que el dashboard gerencial.

    Genera el HTML en una carpeta temporal (`LOCALAPPDATA/FEMAG/recompra`) y lo
    abre en el navegador con ``open()``, sin pedir ruta al usuario.
    """

    def __init__(
        self,
        *,
        service: ManagerialDashboardService | None = None,
        current: ReportPeriod,
        previous: ReportPeriod | None = None,
        product_id: int | None = None,
        salesperson_id: int | None = None,
        product_label: str = "",
        salesperson_label: str = "",
        root_dir: Path | None = None,
        browser_open=None,
    ) -> None:
        import webbrowser

        self.service = service or ManagerialDashboardService()
        self.current = current
        self.previous = previous or current.previous_equivalent()
        self.product_id = product_id
        self.salesperson_id = salesperson_id
        self.product_label = product_label
        self.salesperson_label = salesperson_label
        self.root_dir = root_dir or self._default_root_dir()
        self.browser_open = browser_open or webbrowser.open

    @staticmethod
    def _default_root_dir() -> Path:
        base = os.getenv("LOCALAPPDATA")
        if base:
            return Path(base) / "FEMAG" / "recompra"
        return Path.home() / ".femag" / "recompra"

    @property
    def html_path(self) -> Path:
        return self.root_dir / f"recompra_{self.current.start:%Y%m%d}_{self.current.end:%Y%m%d}.html"

    def generate(self) -> Path:
        self.root_dir.mkdir(parents=True, exist_ok=True)
        doc = _render_repurchase_html(
            self.service,
            self.current,
            self.previous,
            self.product_id,
            self.product_label,
            salesperson_id=self.salesperson_id,
            salesperson_label=self.salesperson_label,
        )
        self.html_path.write_text(doc, encoding="utf-8")
        return self.html_path

    def open(self) -> Path:
        path = self.generate()
        self.browser_open(path.resolve().as_uri())
        return path


def _render_repurchase_html(
    service: ManagerialDashboardService,
    current: ReportPeriod,
    previous: ReportPeriod,
    product_id: int | None,
    product_label: str,
    *,
    salesperson_id: int | None = None,
    salesperson_label: str = "",
) -> str:
    import html as _html

    previous = previous or current.previous_equivalent()
    datasets = build_repurchase_datasets(
        service, current, previous, product_id=product_id, salesperson_id=salesperson_id
    )
    summary = dict(datasets["Resumen"][1])
    rows = service.product_client_repurchase(
        current, previous, product_id=product_id, salesperson_id=salesperson_id
    )
    prev_tn = float(summary.get("tn_anterior") or 0)
    cur_tn = float(summary.get("tn_actual") or 0)
    var = summary.get("variacion_pct")
    var_text = "—" if var is None else f"{var:+.2f}%".replace(".", ",")
    ratio = min(cur_tn / prev_tn, 1.0) if prev_tn > 0 else (1.0 if cur_tn > 0 else 0.0)
    bar_color = "#c43b47" if cur_tn < prev_tn else "#11875d"
    prev_short = f"{previous.start.day}–{previous.end.day} {previous.start:%b}"
    cur_short = f"{current.start.day}–{current.end.day} {current.start:%b}"

    def money(value: float) -> str:
        return f"$ {float(value or 0):,.0f}".replace(",", "X").replace(".", ",").replace("X", ".")

    def table(title: str, items: list[dict]) -> str:
        if not items:
            return f"<h3>{_html.escape(title)}</h3><p>Sin clientes en este grupo.</p>"
        body = "".join(
            f"<tr><td>{_html.escape(r['client_name'])}</td>"
            f"<td class='num'>{r['previous_tonnes']:.3f}</td>"
            f"<td class='num'>{r['current_tonnes']:.3f}</td>"
            f"<td class='num'>{r['diff_tonnes']:+.3f}</td>"
            f"<td>{_html.escape(r.get('salesperson_name') or '—')}</td></tr>"
            for r in items
        )
        return (
            f"<h3>{_html.escape(title)} ({len(items)})</h3>"
            "<table><thead><tr><th>Cliente</th><th class='num'>TN ant.</th>"
            "<th class='num'>TN actual</th><th class='num'>Dif.</th><th>Vendedor</th></tr></thead>"
            f"<tbody>{body}</tbody></table>"
        )

    by_status: dict[str, list[dict]] = {}
    for r in rows:
        by_status.setdefault(r["status"], []).append(r)
    doc = f"""<!doctype html>
<html lang="es"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Recompra por producto · {_html.escape(product_label or 'todos')}</title>
<style>
body{{font-family:Arial,sans-serif;margin:0;background:#0b0d10;color:#e8eaed}}
.shell{{max-width:960px;margin:auto;padding:24px}}
.kpis{{display:flex;gap:24px}}.kpi{{flex:1}}
.kpi .big{{font-size:30px;font-weight:800}}.muted{{color:#9aa0a6}}
.varbar{{background:#3c4043;border-radius:8px;height:12px;margin:10px 0}}
.varbar>div{{height:12px;border-radius:8px;background:{bar_color};width:{ratio * 100:.1f}%}}
.card{{background:#17191d;border:1px solid #2a2d33;border-radius:14px;padding:18px;margin-top:14px}}
table{{width:100%;border-collapse:collapse}}th,td{{padding:8px;border-bottom:1px solid #2a2d33;text-align:left}}
.num{{text-align:right}}h3{{margin:18px 0 8px}}
</style></head><body><div class="shell">
<p class="muted">FEMAG · Análisis de recompra por producto</p>
<h1>{_html.escape(product_label or 'Todos los artículos')}</h1>
<p class="muted">{_html.escape(previous.describe_range())} → {_html.escape(current.describe_range())}{f" · Vendedor: {_html.escape(salesperson_label)}" if salesperson_label else ""}</p>
<div class="kpis">
<div class="kpi"><div class="muted">{_html.escape(prev_short)}</div>
<div class="big">{prev_tn:,.2f} tn</div><div class="muted">{money(summary.get('valorizado_anterior'))}</div></div>
<div class="kpi"><div class="muted">{_html.escape(cur_short)}</div>
<div class="big">{cur_tn:,.2f} tn</div><div class="muted">{money(summary.get('valorizado_actual'))}</div></div>
</div>
<div class="card"><strong>Variación en toneladas</strong>
<span style="float:right;color:{bar_color};font-weight:800">{var_text}</span>
<div class="varbar"><div></div></div>
<p class="muted">Se despacharon {abs(cur_tn - prev_tn):,.2f} toneladas
{'menos' if cur_tn < prev_tn else 'más'} en los mismos {current.days} días.</p>
{_svg_bars(prev_tn, cur_tn, prev_short, cur_short)}
<p class="muted">Volumen despachado del artículo, en toneladas. Conciliación: {summary.get('conciliacion')}</p>
</div>
<div class="card">{table('Clientes sin recompra ⚠ (alerta comercial, no pérdida confirmada)', sorted(by_status.get('sin_recompra', []), key=lambda x: -x['previous_kilos']))}</div>
<div class="card">{table('Compraron menos', sorted(by_status.get('compro_menos', []), key=lambda x: x['diff_kilos']))}</div>
<div class="card">{table('Compraron más', sorted(by_status.get('compro_mas', []), key=lambda x: -x['diff_kilos']))}</div>
<div class="card">{table('Nuevos en el período', sorted(by_status.get('nuevo', []), key=lambda x: -x['current_kilos']))}</div>
</div></body></html>"""
    return doc
