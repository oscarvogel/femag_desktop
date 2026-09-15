from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.reports.managerial_sales_dispatch import (
    ManagerialSalesDispatchService,
    SalesDispatchFilters,
    SalesDispatchReportResult,
)


@dataclass(frozen=True)
class ProductSalesSummary:
    product_id: int
    product_name: str
    quantity: float
    unit: str
    kilos: float
    tonnes: float
    net: float
    vat: float
    total: float
    orders: int


@dataclass(frozen=True)
class OperationalProductSalesResult:
    filters: SalesDispatchFilters
    summary_rows: tuple[ProductSalesSummary, ...]
    detail: SalesDispatchReportResult

    @property
    def total(self) -> float:
        return self.detail.totals.total

    @property
    def tonnes(self) -> float:
        return self.detail.totals.tonnes

    @property
    def products(self) -> int:
        return len(self.summary_rows)

    @property
    def orders(self) -> int:
        return self.detail.totals.orders


class OperationalProductSalesService:
    def __init__(self, *, sales_service: ManagerialSalesDispatchService | None = None) -> None:
        self.sales_service = sales_service or ManagerialSalesDispatchService()

    def report(self, filters: SalesDispatchFilters) -> OperationalProductSalesResult:
        detail = self.sales_service.report(filters)
        buckets: dict[int, dict] = defaultdict(
            lambda: {
                "product_name": "",
                "quantity": 0.0,
                "units": set(),
                "kilos": 0.0,
                "net": 0.0,
                "vat": 0.0,
                "total": 0.0,
                "orders": set(),
            }
        )
        for row in detail.rows:
            bucket = buckets[int(row["product_id"])]
            bucket["product_name"] = row["product_name"] or ""
            bucket["quantity"] += float(row["quantity"] or 0)
            if row.get("unit"):
                bucket["units"].add(str(row["unit"]))
            bucket["kilos"] += float(row["kilos"] or 0)
            bucket["net"] += float(row["net"] or 0)
            bucket["vat"] += float(row["vat"] or 0)
            bucket["total"] += float(row["total"] or 0)
            bucket["orders"].add(int(row["order_id"]))

        summary = []
        for product_id, bucket in buckets.items():
            units = sorted(bucket["units"])
            unit = units[0] if len(units) == 1 else (" / ".join(units) if units else "")
            kilos = round(bucket["kilos"], 3)
            summary.append(
                ProductSalesSummary(
                    product_id=product_id,
                    product_name=bucket["product_name"],
                    quantity=round(bucket["quantity"], 3),
                    unit=unit,
                    kilos=kilos,
                    tonnes=round(kilos / 1000.0, 3),
                    net=round(bucket["net"], 2),
                    vat=round(bucket["vat"], 2),
                    total=round(bucket["total"], 2),
                    orders=len(bucket["orders"]),
                )
            )
        summary.sort(key=lambda row: row.product_name.casefold())
        return OperationalProductSalesResult(
            filters=filters,
            summary_rows=tuple(summary),
            detail=detail,
        )

    def export_xlsx(self, result: OperationalProductSalesResult, destination: str | Path) -> Path:
        path = Path(destination)
        path.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        ws = wb.active
        ws.title = "Resumen por producto"
        headers = ("Producto", "Cantidad", "Unidad", "Kg", "TN", "Neto", "IVA", "Total", "Órdenes")
        ws.append(headers)
        self._style_header(ws)
        for row in result.summary_rows:
            ws.append(
                (
                    row.product_name,
                    row.quantity,
                    row.unit,
                    row.kilos,
                    row.tonnes,
                    row.net,
                    row.vat,
                    row.total,
                    row.orders,
                )
            )
        ws.append(
            (
                "TOTAL",
                "",
                "",
                result.detail.totals.kilos,
                result.detail.totals.tonnes,
                result.detail.totals.net,
                result.detail.totals.vat,
                result.detail.totals.total,
                result.detail.totals.orders,
            )
        )
        total_row = ws.max_row
        for cell in ws[total_row]:
            cell.font = Font(bold=True)
        ws.auto_filter.ref = f"A1:I{max(1, ws.max_row - 1)}"
        ws.freeze_panes = "A2"
        self._format_summary_sheet(ws)

        detail_ws = wb.create_sheet("Detalle de ventas")
        detail_headers = (
            "Fecha",
            "Orden",
            "Cliente",
            "Destino",
            "Producto",
            "Cantidad",
            "Unidad",
            "Kg",
            "TN",
            "Precio neto",
            "Neto",
            "IVA",
            "Total",
            "Estado",
            "Transportista",
        )
        detail_ws.append(detail_headers)
        self._style_header(detail_ws)
        for row in result.detail.rows:
            detail_ws.append(
                (
                    row["date"],
                    row["order_number"],
                    row["client_name"],
                    row["destination"],
                    row["product_name"],
                    row["quantity"],
                    row["unit"],
                    row["kilos"],
                    row["tonnes"],
                    row["unit_net_price"],
                    row["net"],
                    row["vat"],
                    row["total"],
                    row["status"],
                    row["carrier_name"],
                )
            )
        detail_ws.auto_filter.ref = f"A1:O{max(1, detail_ws.max_row)}"
        detail_ws.freeze_panes = "A2"
        self._format_detail_sheet(detail_ws)

        wb.save(path)
        return path

    @staticmethod
    def _style_header(ws) -> None:
        for cell in ws[1]:
            cell.font = Font(bold=True)

    @staticmethod
    def _format_summary_sheet(ws) -> None:
        widths = (34, 14, 12, 14, 12, 16, 16, 16, 10)
        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        for row in range(2, ws.max_row + 1):
            for col in (2, 4, 5):
                ws.cell(row, col).number_format = '#,##0.000'
            for col in (6, 7, 8):
                ws.cell(row, col).number_format = '$ #,##0.00'

    @staticmethod
    def _format_detail_sheet(ws) -> None:
        widths = (13, 11, 28, 28, 34, 14, 12, 14, 12, 16, 16, 16, 16, 14, 28)
        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        for row in range(2, ws.max_row + 1):
            ws.cell(row, 1).number_format = 'dd/mm/yyyy'
            for col in (6, 8, 9):
                ws.cell(row, col).number_format = '#,##0.000'
            for col in (10, 11, 12, 13):
                ws.cell(row, col).number_format = '$ #,##0.00'
