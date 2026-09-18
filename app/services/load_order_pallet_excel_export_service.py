from __future__ import annotations

from collections import OrderedDict
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image as PILImage
from reportlab.graphics import renderPM

from app.models.load_orders import LoadOrder
from app.services.audit_service import AuditService
from app.services.qr_load_order_print_service import ConsolidatedLoadOrderPrintService


class LoadOrderPalletExcelExportService:
    """Crea un Excel editable con la misma estructura visual de la orden PDF."""

    HEADERS = ("Producto / detalle", "Cantidad total", "Pallet", "Lote", "Elab.")
    COL_WIDTHS = (58, 20, 16, 16, 16)

    def __init__(self, current_user: str, audit_service: AuditService | None = None) -> None:
        self.current_user = current_user
        self.audit_service = audit_service or AuditService()

    def export(self, order: LoadOrder, output_dir: str | Path) -> Path:
        order = LoadOrder.get_by_id(order.id)
        destination = Path(output_dir)
        destination.mkdir(parents=True, exist_ok=True)
        target = destination / f"orden_carga_{order.order_number}_armado_pallets.xlsx"

        printer = ConsolidatedLoadOrderPrintService(current_user=self.current_user)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Orden de despacho"
        sheet.sheet_view.showGridLines = False

        for index, width in enumerate(self.COL_WIDTHS, start=1):
            sheet.column_dimensions[chr(64 + index)].width = width

        row = 1
        row = self._write_header(sheet, row, order, printer)
        row = self._write_client_section(sheet, row, order, printer)
        row = self._write_detail_section(sheet, row, order, printer)
        row = self._write_transport_section(sheet, row, order, printer)
        row = self._write_footer(sheet, row, order)

        sheet.freeze_panes = "A1"
        sheet.page_setup.orientation = "portrait"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_options.horizontalCentered = True
        sheet.print_area = f"A1:E{row}"
        sheet.page_margins.left = 0.3
        sheet.page_margins.right = 0.3
        sheet.page_margins.top = 0.4
        sheet.page_margins.bottom = 0.4

        workbook.save(target)
        self.audit_service.record(
            user=self.current_user,
            module="Ordenes de carga",
            action="exportar_excel_pallets",
            record_ref=f"LoadOrder:{order.id}",
            new_value={"file_path": str(target), "order_number": order.order_number},
        )
        return target

    def _write_header(self, sheet, row: int, order: LoadOrder, printer) -> int:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        title = sheet.cell(row, 1, "ORDEN DE DESPACHO DE FECULA DE MANDIOCA")
        title.font = Font(bold=True, size=14)
        title.alignment = Alignment(horizontal="left", vertical="center")
        sheet.row_dimensions[row].height = 24
        row += 1

        sheet.cell(row, 1, f"Nro.: {order.order_number:04d}").font = Font(bold=True)
        sheet.cell(row, 2, f"Fecha: {order.date:%d/%m/%Y}").font = Font(bold=True)
        sheet.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
        sheet.cell(row, 3, f"Estado: {order.status}").font = Font(bold=True)
        row += 2

        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        qr_title = sheet.cell(row, 1, "QR de la orden")
        qr_title.font = Font(bold=True, size=10)
        row += 1

        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        sheet.cell(
            row,
            1,
            "Identificador para carga de lote y fecha de elaboración.",
        ).font = Font(size=8)

        self._add_qr_image(sheet, order, printer, anchor=f"E3")
        sheet.row_dimensions[row].height = 24
        row += 3
        return row

    def _write_client_section(self, sheet, row: int, order: LoadOrder, printer) -> int:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        section = sheet.cell(row, 1, "1. DATOS DEL CLIENTE")
        section.font = Font(bold=True, size=11)
        row += 1

        sheet.cell(row, 1, "DATOS DEL CLIENTE:").font = Font(size=8)
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
        sheet.cell(row, 2, printer._client_label(order)).font = Font(size=8)
        row += 1

        sheet.cell(row, 1, "DESTINO:").font = Font(size=8)
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
        destination = sheet.cell(row, 2, printer._destination_label(order))
        destination.font = Font(size=8)
        destination.alignment = Alignment(wrap_text=True, vertical="top")
        row += 3
        return row

    def _write_detail_section(self, sheet, row: int, order: LoadOrder, printer) -> int:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        section = sheet.cell(row, 1, "2. DETALLE DEL PRODUCTO A DESPACHAR")
        section.font = Font(bold=True, size=11)
        row += 1

        blocks = printer._detail_blocks(order)
        first_detail_row = row
        for block in blocks:
            row = self._write_destination_block(sheet, row, block, printer)

        row += 1
        row = self._write_totals(sheet, row, order, blocks, printer)

        self._apply_grid(sheet, first_detail_row, row - 1)
        row += 3
        return row

    def _write_destination_block(self, sheet, row: int, block: dict[str, object], printer) -> int:
        header_fill = PatternFill("solid", fgColor="F2F2F2")
        destination_fill = PatternFill("solid", fgColor="E7E6E6")

        for column, header in enumerate(self.HEADERS, start=1):
            cell = sheet.cell(row, column, header)
            cell.font = Font(bold=True, size=8)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        row += 1

        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        destination = sheet.cell(row, 1, str(block.get("destination") or "-"))
        destination.font = Font(bold=True, size=8)
        destination.fill = destination_fill
        destination.alignment = Alignment(vertical="center", wrap_text=True)
        row += 1

        sub_blocks = [
            *block.get("pallet_blocks", []),
            *([block.get("loose_block")] if block.get("loose_block") else []),
            *([block.get("unassigned_block")] if block.get("unassigned_block") else []),
        ]
        index = 0
        has_rows = False
        while index < len(sub_blocks):
            sub_block = sub_blocks[index]
            physical_rows = list(sub_block.get("rows", []))
            if not physical_rows:
                index += 1
                continue

            group_key = printer._single_product_pallet_key(sub_block)
            if group_key is not None:
                group_end = index + 1
                while (
                    group_end < len(sub_blocks)
                    and printer._single_product_pallet_key(sub_blocks[group_end]) == group_key
                ):
                    group_end += 1
                item = physical_rows[0]
                self._append_detail_row(
                    sheet,
                    row,
                    printer,
                    item,
                    printer._pallet_label(group_end - index),
                )
                row += 1
                index = group_end
                has_rows = True
                continue

            start_row = row
            label = str(sub_block.get("label") or "-")
            for item_index, item in enumerate(physical_rows):
                self._append_detail_row(
                    sheet,
                    row,
                    printer,
                    item,
                    label if item_index == 0 else "",
                )
                row += 1
            if len(physical_rows) > 1:
                sheet.merge_cells(
                    start_row=start_row,
                    start_column=3,
                    end_row=row - 1,
                    end_column=3,
                )
                sheet.cell(start_row, 3).alignment = Alignment(
                    horizontal="center", vertical="center"
                )
            index += 1
            has_rows = True

        if not has_rows:
            for column in range(1, 6):
                sheet.cell(row, column, "-")
            row += 1
        return row

    def _write_totals(self, sheet, row: int, order: LoadOrder, blocks: list[dict[str, object]], printer) -> int:
        totals: OrderedDict[str, float] = OrderedDict()
        for block in blocks:
            sub_blocks = [
                *block.get("pallet_blocks", []),
                *([block.get("loose_block")] if block.get("loose_block") else []),
                *([block.get("unassigned_block")] if block.get("unassigned_block") else []),
            ]
            for sub_block in sub_blocks:
                for item in sub_block.get("rows", []):
                    product = str(item.get("product") or "-")
                    totals[product] = totals.get(product, 0.0) + float(item.get("quantity") or 0)

        for product, quantity in totals.items():
            sheet.cell(row, 1, "")
            sheet.cell(row, 2, printer._quantity_with_unit(quantity, None))
            sheet.cell(row, 3, "")
            sheet.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
            sheet.cell(row, 3, product)
            sheet.cell(row, 2).font = Font(bold=True, size=8)
            sheet.cell(row, 3).font = Font(size=8)
            sheet.cell(row, 2).alignment = Alignment(horizontal="center", vertical="center")
            row += 1

        sheet.cell(row, 1, "TOTALES").font = Font(bold=True, size=8)
        sheet.cell(
            row,
            2,
            f"{printer._used_pallet_total(order)} pallet"
            + ("s" if printer._used_pallet_total(order) != 1 else ""),
        ).font = Font(bold=True, size=8)
        sheet.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
        sheet.cell(row, 1).alignment = Alignment(horizontal="center", vertical="center")
        sheet.cell(row, 2).alignment = Alignment(horizontal="center", vertical="center")
        return row + 1

    def _write_transport_section(self, sheet, row: int, order: LoadOrder, printer) -> int:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        section = sheet.cell(row, 1, "3. DATOS DEL TRANSPORTE")
        section.font = Font(bold=True, size=11)
        row += 1

        rows = [
            ("Empresa de transporte:", order.carrier.name),
            ("Dominio del vehiculo:", order.truck.domain),
            ("Nombre del chofer:", order.driver.name),
            ("Dominio semi/acoplado:", printer._snapshot_trailer_domain(order)),
            ("TOTAL MERCADERIA:", printer._kg_text(printer._merchandise_kg_total(order))),
            ("Vehiculo limpio y apto:", "Si / No"),
        ]
        start = row
        for label, value in rows:
            sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
            sheet.cell(row, 1, label).font = Font(bold=True, size=8)
            sheet.cell(row, 2, value).font = Font(
                bold=(label == "TOTAL MERCADERIA:"), size=8
            )
            sheet.cell(row, 1).alignment = Alignment(horizontal="center", vertical="center")
            sheet.cell(row, 2).alignment = Alignment(vertical="center")
            row += 1
        self._apply_grid(sheet, start, row - 1)
        row += 3
        return row

    def _write_footer(self, sheet, row: int, order: LoadOrder) -> int:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row + 1, end_column=5)
        observations = sheet.cell(
            row,
            1,
            f"Observaciones: {order.observations or '-'}",
        )
        observations.font = Font(size=9)
        observations.alignment = Alignment(wrap_text=True, vertical="top")
        row += 4

        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        sheet.cell(
            row,
            1,
            "Firma del encargado de carga: __________________________",
        ).font = Font(size=9)
        return row + 1

    @staticmethod
    def _append_detail_row(sheet, row_number: int, printer, item: dict, pallet: str) -> None:
        sheet.cell(row_number, 1, item.get("product") or "-")
        sheet.cell(
            row_number,
            2,
            printer._quantity_with_unit(item.get("quantity"), item.get("unit")),
        )
        sheet.cell(row_number, 3, pallet)
        sheet.cell(
            row_number,
            4,
            item.get("lote") if printer._optional_operational_value(item.get("lote")) else "",
        )
        sheet.cell(
            row_number,
            5,
            item.get("elab") if printer._optional_operational_value(item.get("elab")) else "",
        )
        for column in range(1, 6):
            sheet.cell(row_number, column).font = Font(size=8)
            sheet.cell(row_number, column).alignment = Alignment(
                vertical="center", wrap_text=True
            )
        for column in (2, 3):
            sheet.cell(row_number, column).alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            sheet.cell(row_number, column).font = Font(bold=True, size=9)

    @staticmethod
    def _apply_grid(sheet, start_row: int, end_row: int) -> None:
        thin = Side(style="thin", color="404040")
        for row in sheet.iter_rows(
            min_row=start_row,
            max_row=end_row,
            min_col=1,
            max_col=5,
        ):
            for cell in row:
                cell.border = Border(
                    left=thin,
                    right=thin,
                    top=thin,
                    bottom=thin,
                )

    @staticmethod
    def _add_qr_image(sheet, order: LoadOrder, printer) -> None:
        try:
            png = renderPM.drawToString(printer._qr_drawing(order), fmt="PNG")
            image = PILImage.open(BytesIO(png))
            image.load()
            excel_image = ExcelImage(image)
            excel_image.width = 78
            excel_image.height = 78
            sheet.add_image(excel_image, "E3")
        except Exception:
            # El Excel debe seguir generándose aun si el backend gráfico de
            # ReportLab no está disponible en una instalación puntual.
            return
