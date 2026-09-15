from datetime import date

from openpyxl import load_workbook

from app.reports.managerial_sales_dispatch import (
    SalesDispatchFilters,
    SalesDispatchReportResult,
    SalesDispatchTotals,
)
from app.reports.operational_product_sales import OperationalProductSalesService


class FakeSalesService:
    def report(self, filters):
        rows = (
            {
                "line_id": 1,
                "date": date(2026, 9, 1),
                "order_id": 10,
                "order_number": 1001,
                "status": "Cerrada",
                "carrier_id": 1,
                "carrier_name": "Transportista A",
                "client_id": 1,
                "client_name": "Cliente A",
                "destination_id": 1,
                "destination": "Posadas",
                "product_id": 7,
                "product_name": "Fécula A",
                "quantity": 10.0,
                "unit": "bolsa",
                "kilos": 250.0,
                "tonnes": 0.25,
                "unit_net_price": 1000.0,
                "net": 10000.0,
                "vat": 2100.0,
                "total": 12100.0,
            },
            {
                "line_id": 2,
                "date": date(2026, 9, 2),
                "order_id": 11,
                "order_number": 1002,
                "status": "Emitida",
                "carrier_id": 1,
                "carrier_name": "Transportista A",
                "client_id": 2,
                "client_name": "Cliente B",
                "destination_id": 2,
                "destination": "Eldorado",
                "product_id": 7,
                "product_name": "Fécula A",
                "quantity": 20.0,
                "unit": "bolsa",
                "kilos": 500.0,
                "tonnes": 0.5,
                "unit_net_price": 1000.0,
                "net": 20000.0,
                "vat": 4200.0,
                "total": 24200.0,
            },
            {
                "line_id": 3,
                "date": date(2026, 9, 2),
                "order_id": 11,
                "order_number": 1002,
                "status": "Emitida",
                "carrier_id": 1,
                "carrier_name": "Transportista A",
                "client_id": 2,
                "client_name": "Cliente B",
                "destination_id": 2,
                "destination": "Eldorado",
                "product_id": 8,
                "product_name": "Fécula B",
                "quantity": 5.0,
                "unit": "bolsa",
                "kilos": 100.0,
                "tonnes": 0.1,
                "unit_net_price": 2000.0,
                "net": 10000.0,
                "vat": 2100.0,
                "total": 12100.0,
            },
        )
        return SalesDispatchReportResult(
            filters=filters,
            rows=rows,
            totals=SalesDispatchTotals(
                net=40000.0,
                vat=8400.0,
                total=48400.0,
                kilos=850.0,
                tonnes=0.85,
                orders=2,
                lines=3,
            ),
        )


def test_operational_sales_groups_by_product():
    filters = SalesDispatchFilters(date(2026, 9, 1), date(2026, 9, 30))
    result = OperationalProductSalesService(sales_service=FakeSalesService()).report(filters)

    assert result.products == 2
    assert result.orders == 2
    assert result.total == 48400.0
    assert result.tonnes == 0.85

    first = result.summary_rows[0]
    assert first.product_name == "Fécula A"
    assert first.quantity == 30.0
    assert first.kilos == 750.0
    assert first.tonnes == 0.75
    assert first.net == 30000.0
    assert first.vat == 6300.0
    assert first.total == 36300.0
    assert first.orders == 2


def test_operational_sales_exports_two_sheet_xlsx(tmp_path):
    filters = SalesDispatchFilters(date(2026, 9, 1), date(2026, 9, 30))
    service = OperationalProductSalesService(sales_service=FakeSalesService())
    result = service.report(filters)

    target = service.export_xlsx(result, tmp_path / "ventas.xlsx")

    assert target.exists()
    workbook = load_workbook(target, data_only=True)
    assert workbook.sheetnames == ["Resumen por producto", "Detalle de ventas"]

    summary = workbook["Resumen por producto"]
    detail = workbook["Detalle de ventas"]

    assert summary["A1"].value == "Producto"
    assert summary["A2"].value == "Fécula A"
    assert summary["H2"].value == 36300
    assert summary.cell(summary.max_row, 1).value == "TOTAL"
    assert summary.cell(summary.max_row, 8).value == 48400
    assert detail["A1"].value == "Fecha"
    assert detail["B2"].value == 1001
    assert detail["E2"].value == "Fécula A"
    assert detail["M2"].value == 12100
