import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")


def _user(username, profile_name="Administrador"):
    from app.models.security import User, UserProfile

    profile, _ = UserProfile.get_or_create(name=profile_name)
    return User.create(username=username, password_hash="x", profile=profile, active=True)


def test_export_clients_xlsx_includes_lista_and_contact_data(db, tmp_path):
    from openpyxl import load_workbook

    from app.models.masters import Client, Salesperson
    from app.reports.client_list_export import HEADERS, export_clients_xlsx

    vendor = Salesperson.create(name="Vendedor 704", phone="3764000000")
    Client.create(
        name="Cliente Lista Uno",
        cuit="30704000001",
        iva_condition="RI",
        phone="3764111111",
        email="uno@example.com",
        contact="Contacto Uno",
        salesperson=vendor,
        lista_precios=1,
    )
    Client.create(
        name="Cliente Lista Dos",
        cuit="30704000002",
        iva_condition="RI",
        phone="3764222222",
        email="dos@example.com",
        contact="Contacto Dos",
        lista_precios=2,
    )

    destination = tmp_path / "clientes_704.xlsx"
    exported = export_clients_xlsx(list(Client.select().order_by(Client.name)), destination)

    wb = load_workbook(exported)
    ws = wb.active
    assert [cell.value for cell in ws[1]] == list(HEADERS)
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2
    by_name = {row[0]: row for row in rows}
    assert by_name["Cliente Lista Uno"][3] == "3764111111"
    assert by_name["Cliente Lista Uno"][6] == "Vendedor 704"
    assert by_name["Cliente Lista Uno"][7] == 1
    assert by_name["Cliente Lista Dos"][7] == 2
    assert by_name["Cliente Lista Dos"][5] == "Contacto Dos"


def test_client_export_prefers_primary_contact_email(db, tmp_path):
    from openpyxl import load_workbook

    from app.models.masters import Client, ClientEmail
    from app.reports.client_list_export import export_clients_xlsx

    client = Client.create(
        name="Cliente Mails 704",
        cuit="30704000003",
        iva_condition="RI",
        phone="3764333333",
        email="base@example.com",
        lista_precios=3,
    )
    ClientEmail.create(client=client, email="secundario@example.com", is_primary=False)
    ClientEmail.create(client=client, email="principal@example.com", is_primary=True)

    destination = tmp_path / "clientes_mails_704.xlsx"
    exported = export_clients_xlsx([Client.get_by_id(client.id)], destination)

    wb = load_workbook(exported)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert rows[0][4] == "principal@example.com"


def test_clients_page_has_export_excel_button(db):
    from PyQt5.QtWidgets import QApplication, QPushButton

    from app.services.permission_service import PermissionService
    from app.ui.master_abm import build_client_abm_page

    PermissionService().seed_defaults()
    app = QApplication.instance() or QApplication([])
    admin = _user("admin-clientes-excel-704")
    page = build_client_abm_page(
        user=admin,
        current_user=admin.username,
        view_ledger_callback=lambda _client: None,
    )
    app.processEvents()

    button = page.findChild(QPushButton, "exportClientsExcelButton")
    assert button is not None
    assert button.text() == "Exportar a Excel"
