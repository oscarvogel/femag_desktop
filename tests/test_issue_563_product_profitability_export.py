import csv
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from app.reports.product_profitability import ProductProfitabilityService
from app.reports.product_profitability_export import export_profitability_csv, export_profitability_xlsx


def test_profitability_exports_empty_snapshot(db, tmp_path):
    snap=ProductProfitabilityService().snapshot(date.today(),date.today())
    xlsx=export_profitability_xlsx(snap,tmp_path/"profit.xlsx",start=date.today(),end=date.today())
    csv_path=export_profitability_csv(snap,tmp_path/"profit.csv")
    assert xlsx.exists() and csv_path.exists()
    wb=load_workbook(xlsx,read_only=True)
    assert wb.sheetnames==["Resumen por producto","Detalle por operacion"]
    with csv_path.open(encoding="utf-8-sig") as handle:
        rows=list(csv.reader(handle,delimiter=";"))
    assert rows[0][0:4]==["Fecha","Orden","Cliente","Producto"]
