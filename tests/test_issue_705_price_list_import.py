"""Importación de lista de precios y teléfono desde Excel (#705)."""

import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")


def _user(username, profile_name="Administrador"):
    from app.models.security import User, UserProfile

    profile, _ = UserProfile.get_or_create(name=profile_name)
    return User.create(username=username, password_hash="x", profile=profile, active=True)


def _xlsx(tmp_path, filas, nombre="clientes_705.xlsx"):
    """Arma un Excel con el mismo formato que exporta la app."""
    from openpyxl import Workbook

    from app.reports.client_list_export import HEADERS

    libro = Workbook()
    hoja = libro.active
    hoja.append(list(HEADERS))
    for fila in filas:
        hoja.append(fila)
    destino = tmp_path / nombre
    libro.save(destino)
    return destino


def _fila(nombre, cuit, lista, telefono="", email="", contacto="", iva="RI", estado="Activo"):
    return [nombre, cuit, iva, telefono, email, contacto, "Sin asignar", lista, estado]


def _cliente(nombre="Cliente 705", cuit="30705000001", lista=1, phone=None, iva="RI", email="", contact=""):
    from app.models.masters import Client

    return Client.create(
        name=nombre,
        cuit=cuit,
        iva_condition=iva,
        phone=phone,
        email=email or None,
        contact=contact or None,
        lista_precios=lista,
    )


def _importador():
    from app.services.client_price_list_import_service import ClientPriceListImporter

    return ClientPriceListImporter(current_user="admin-705")


def test_import_cambia_solo_la_lista_y_no_toca_el_resto(db, tmp_path):
    cliente = _cliente(nombre="Cliente Uno", cuit="30705000001", lista=1, phone="3764111111", iva="RI", email="a@x.com", contact="Ana")
    archivo = _xlsx(tmp_path, [_fila("Cliente Uno", "30705000001", 2, telefono="3764999999")])

    vista = _importador().construir_vista_previa(archivo)
    assert len(vista.cambios_lista) == 1
    assert not vista.cambios_telefono

    resultado = _importador().aplicar(vista)
    assert resultado.errores == []
    assert resultado.cambios_lista == 1

    from app.models.masters import Client

    guardado = Client.get_by_id(cliente.id)
    assert guardado.lista_precios == 2
    assert guardado.phone == "3764111111"
    assert guardado.name == "Cliente Uno"
    assert guardado.cuit == "30705000001"
    assert guardado.email == "a@x.com"
    assert guardado.contact == "Ana"
    assert guardado.iva_condition == "RI"


def test_telefono_se_completa_solo_si_esta_vacio(db, tmp_path):
    cliente = _cliente(nombre="Sin Telefono", cuit="30705000002", lista=1, phone=None)
    archivo = _xlsx(tmp_path, [_fila("Sin Telefono", "30705000002", 1, telefono="3758 650849")])

    vista = _importador().construir_vista_previa(archivo)
    assert len(vista.cambios_telefono) == 1
    assert not vista.cambios_lista

    _importador().aplicar(vista)

    from app.models.masters import Client

    assert Client.get_by_id(cliente.id).phone == "3758 650849"


def test_telefono_existente_no_se_pisa(db, tmp_path):
    cliente = _cliente(nombre="Con Telefono", cuit="30705000003", lista=1, phone="3764111111")
    archivo = _xlsx(tmp_path, [_fila("Con Telefono", "30705000003", 1, telefono="011 34070557")])

    vista = _importador().construir_vista_previa(archivo)
    assert not vista.cambios_telefono
    assert any(diferencia.campo == "teléfono" for diferencia in vista.diferencias)

    _importador().aplicar(vista)

    from app.models.masters import Client

    assert Client.get_by_id(cliente.id).phone == "3764111111"


def test_telefono_vacio_en_el_archivo_no_borra_el_de_la_base(db, tmp_path):
    cliente = _cliente(nombre="Con Telefono", cuit="30705000004", lista=1, phone="3764111111")
    archivo = _xlsx(tmp_path, [_fila("Con Telefono", "30705000004", 2)])

    vista = _importador().construir_vista_previa(archivo)
    _importador().aplicar(vista)

    from app.models.masters import Client

    guardado = Client.get_by_id(cliente.id)
    assert guardado.phone == "3764111111"
    assert guardado.lista_precios == 2


def test_match_por_cuit_con_separadores(db, tmp_path):
    cliente = _cliente(nombre="Con Guiones", cuit="20- 29646152-1", lista=2)
    archivo = _xlsx(tmp_path, [_fila("Con Guiones", "20296461521", 1)])

    vista = _importador().construir_vista_previa(archivo)
    assert len(vista.cambios_lista) == 1
    assert vista.cambios_lista[0].client_id == cliente.id


def test_match_por_nombre_cuando_el_cuit_no_sirve(db, tmp_path):
    cliente = _cliente(nombre="Sin Cuit Utilizable", cuit="00000000000", lista=1)
    archivo = _xlsx(tmp_path, [_fila("Sin Cuit Utilizable", "", 2)])

    vista = _importador().construir_vista_previa(archivo)
    assert len(vista.cambios_lista) == 1
    assert vista.cambios_lista[0].client_id == cliente.id


def test_cuit_duplicado_en_el_archivo_con_listas_distintas_no_se_aplica(db, tmp_path):
    _cliente(nombre="Duplicado Archivo", cuit="30705000005", lista=1)
    archivo = _xlsx(
        tmp_path,
        [
            _fila("Duplicado Archivo", "30705000005", 2),
            _fila("Duplicado Archivo", "30705000005", 1),
        ],
    )

    vista = _importador().construir_vista_previa(archivo)
    assert not vista.cambios
    assert vista.hay_conflictos
    assert vista.conflictos[0].tipo == "cuit_duplicado_archivo"


def test_cuit_ambiguo_en_la_base_no_se_aplica(db, tmp_path):
    _cliente(nombre="Duplicado Base Uno", cuit="30705000006", lista=1)
    _cliente(nombre="Duplicado Base Dos", cuit="30 705000006", lista=1)
    archivo = _xlsx(tmp_path, [_fila("Duplicado Base Uno", "30705000006", 3)])

    vista = _importador().construir_vista_previa(archivo)
    assert not vista.cambios
    assert vista.hay_conflictos
    assert vista.conflictos[0].tipo == "cuit_ambiguo_base"


def test_duplicado_mismo_cliente_en_la_base_se_unifica(db, tmp_path):
    _cliente(nombre="Gimenez Prueba", cuit="30705000007", lista=1)
    _cliente(nombre="Gimenez Prueba", cuit="30-70500000-7", lista=1)
    archivo = _xlsx(tmp_path, [_fila("Gimenez Prueba", "30705000007", 2)])

    vista = _importador().construir_vista_previa(archivo)
    assert len(vista.cambios_lista) == 2
    assert vista.conflictos[0].tipo == "duplicado_unificado"


def test_duplicado_mismo_cliente_con_listas_distintas_no_bloquea_el_cambio(db, tmp_path):
    """#761: el mismo cliente con el CUIT en dos formatos y listas distintas.

    En la base real el CUIT se guardó en dos formatos (el legacy sin guiones y
    la carga manual con guiones) y el UNIQUE es sobre el texto, asi que
    conviven dos filas del mismo cliente con listas distintas. La lista del
    archivo tiene que llegar a la fila que quedo en la lista vieja.
    """
    legacy = _cliente(nombre="Gimenez 761", cuit="30705000016", lista=1)
    manual = _cliente(nombre="Gimenez 761", cuit="30-7050000-16", lista=2)
    archivo = _xlsx(tmp_path, [_fila("Gimenez 761", "30705000016", 2)])

    vista = _importador().construir_vista_previa(archivo)

    assert len(vista.cambios_lista) == 1
    assert vista.cambios_lista[0].client_id == legacy.id
    assert vista.cambios_lista[0].lista_anterior == 1
    assert vista.cambios_lista[0].lista_nueva == 2
    # La fila que ya estaba en la lista 2 no cambia, pero tampoco bloquea.
    assert manual.id not in [cambio.client_id for cambio in vista.cambios_lista]
    # Queda avisado de que hay que unificar las filas, pero no es un bloqueo.
    assert vista.hay_conflictos
    assert vista.conflictos[0].tipo == "duplicado_unificado"


def test_duplicado_mismo_cliente_con_listas_distintas_se_aplica_a_todas_las_filas(db, tmp_path):
    from app.models.masters import Client

    legacy = _cliente(nombre="Unificado 761", cuit="30705000017", lista=1, phone=None)
    manual = _cliente(nombre="Unificado 761", cuit="30-7050000-17", lista=2)
    archivo = _xlsx(
        tmp_path, [_fila("Unificado 761", "30705000017", 2, telefono="011 3116 7114")]
    )

    resultado = _importador().aplicar(_importador().construir_vista_previa(archivo))

    assert resultado.cambios_lista == 1
    assert Client.get_by_id(legacy.id).lista_precios == 2
    assert Client.get_by_id(manual.id).lista_precios == 2
    # El telefono falta en la fila legacy: se completa solo en esa.
    assert Client.get_by_id(legacy.id).phone == "011 3116 7114"
    # No se borra ninguna fila: unificar las filas del legacy es otro trabajo.
    assert Client.select().where(Client.name == "Unificado 761").count() == 2


def test_dos_clientes_distintos_con_el_mismo_cuit_no_se_aplica_aunque_difieran_las_listas(db, tmp_path):
    """Nombres distintos: son clientes distintos, no se aplica nada.

    Antes este caso se reportaba como "varias filas del cliente con listas
    distintas", que hacia creer que era el mismo cliente.
    """
    _cliente(nombre="Empresa Uno 761", cuit="30705000018", lista=1)
    _cliente(nombre="Empresa Dos 761", cuit="30-7050000-18", lista=3)
    archivo = _xlsx(tmp_path, [_fila("Empresa Uno 761", "30705000018", 2)])

    vista = _importador().construir_vista_previa(archivo)

    assert not vista.cambios
    assert vista.hay_conflictos
    assert vista.conflictos[0].tipo == "cuit_ambiguo_base"
    assert "más de un cliente" in vista.conflictos[0].detalle


def test_lista_invalida_no_se_aplica(db, tmp_path):
    _cliente(nombre="Lista Rara", cuit="30705000008", lista=1)
    archivo = _xlsx(tmp_path, [_fila("Lista Rara", "30705000008", 7)])

    vista = _importador().construir_vista_previa(archivo)
    assert not vista.cambios
    assert len(vista.invalidos) == 1


def test_no_crea_clientes_que_no_estan_en_la_base(db, tmp_path):
    from app.models.masters import Client

    archivo = _xlsx(tmp_path, [_fila("Cliente Nuevo", "30709999999", 2)])

    vista = _importador().construir_vista_previa(archivo)
    assert not vista.cambios
    assert len(vista.no_encontrados) == 1

    _importador().aplicar(vista)
    assert Client.select().where(Client.cuit == "30709999999").count() == 0


def test_clientes_que_el_archivo_no_menciona_no_se_tocan(db, tmp_path):
    cliente = _cliente(nombre="No Mencionado", cuit="30705000009", lista=2)
    archivo = _xlsx(tmp_path, [_fila("Otro", "30705000010", 1)])

    vista = _importador().construir_vista_previa(archivo)
    # "Otro" no está en la base: queda como no encontrado. El que sí está y no lo
    # menciona el archivo queda como faltante (no se toca).
    assert len(vista.faltantes) == 1
    assert len(vista.no_encontrados) == 1

    _importador().aplicar(vista)

    from app.models.masters import Client

    assert Client.get_by_id(cliente.id).lista_precios == 2


def test_import_es_idempotente(db, tmp_path):
    _cliente(nombre="Idempotente", cuit="30705000011", lista=1)
    archivo = _xlsx(tmp_path, [_fila("Idempotente", "30705000011", 3)])

    primera = _importador().aplicar(_importador().construir_vista_previa(archivo))
    segunda = _importador().aplicar(_importador().construir_vista_previa(archivo))

    assert primera.cambios_lista == 1
    assert segunda.cambios_lista == 0
    assert segunda.cambios_telefono == 0


def test_la_vista_previa_no_escribe_nada(db, tmp_path):
    from app.models.masters import Client

    cliente = _cliente(nombre="Solo Vista Previa", cuit="30705000012", lista=1)
    archivo = _xlsx(tmp_path, [_fila("Solo Vista Previa", "30705000012", 4)])

    _importador().construir_vista_previa(archivo)

    assert Client.get_by_id(cliente.id).lista_precios == 1


def test_se_guarda_respaldo_antes_de_aplicar(db, tmp_path):
    _cliente(nombre="Con Respaldo", cuit="30705000013", lista=1)
    archivo = _xlsx(tmp_path, [_fila("Con Respaldo", "30705000013", 2)])

    resultado = _importador().aplicar(_importador().construir_vista_previa(archivo))

    assert resultado.respaldo
    assert os.path.isfile(resultado.respaldo)


def test_el_respaldo_tiene_el_estado_anterior(db, tmp_path):
    from openpyxl import load_workbook

    _cliente(nombre="Con Respaldo", cuit="30705000013", lista=1, phone=None)
    archivo = _xlsx(tmp_path, [_fila("Con Respaldo", "30705000013", 2, telefono="3758 650849")])

    resultado = _importador().aplicar(_importador().construir_vista_previa(archivo))

    libro = load_workbook(resultado.respaldo)
    filas = list(libro.active.iter_rows(values_only=True))
    encabezado = filas[0]
    datos = {fila[2]: fila for fila in filas[1:]}
    assert "CUIT" in encabezado
    assert datos["30705000013"][4] == 1
    assert datos["30705000013"][3] in ("", None)


def test_cada_cambio_deja_actualizado_el_momento(db, tmp_path):
    from app.models.masters import Client

    cliente = _cliente(nombre="Con Marca De Tiempo", cuit="30705000015", lista=1)
    antes = str(Client.get_by_id(cliente.id).updated_at)
    archivo = _xlsx(tmp_path, [_fila("Con Marca De Tiempo", "30705000015", 2)])

    _importador().aplicar(_importador().construir_vista_previa(archivo))

    assert str(Client.get_by_id(cliente.id).updated_at) != antes


def test_cada_cambio_queda_auditado(db, tmp_path):
    from app.models.audit import AuditLog

    cliente = _cliente(nombre="Auditado", cuit="30705000014", lista=1, phone=None)
    archivo = _xlsx(
        tmp_path,
        [_fila("Auditado", "30705000014", 2, telefono="3758 650849")],
    )

    _importador().aplicar(_importador().construir_vista_previa(archivo))

    acciones = {row.action for row in AuditLog.select().where(AuditLog.record_ref == f"Client:{cliente.id}")}
    assert "actualizar lista de precios" in acciones
    assert "completar teléfono" in acciones


def test_clientes_page_has_import_button(db):
    from PyQt5.QtWidgets import QApplication, QPushButton

    from app.services.permission_service import PermissionService
    from app.ui.master_abm import build_client_abm_page

    PermissionService().seed_defaults()
    app = QApplication.instance() or QApplication([])
    admin = _user("admin-import-705")
    page = build_client_abm_page(
        user=admin,
        current_user=admin.username,
        view_ledger_callback=lambda _client: None,
    )
    app.processEvents()

    button = page.findChild(QPushButton, "importPriceListExcelButton")
    assert button is not None
    assert button.text() == "Importar lista"
