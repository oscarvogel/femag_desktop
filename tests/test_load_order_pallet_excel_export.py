from openpyxl import load_workbook


def test_pallet_excel_matches_print_layout_and_keeps_mixed_pallet_cells_merged(db, tmp_path):
    from app.models.audit import AuditLog
    from app.models.masters import Carrier, Client, ClientAddress, Driver, PalletType, Product, Truck
    from app.services.load_order_pallet_excel_export_service import LoadOrderPalletExcelExportService
    from app.services.load_order_service import LoadOrderService

    client = Client.create(name="Cliente Excel", cuit="30700000123", iva_condition="RI")
    address = ClientAddress.create(
        client=client,
        address_type="entrega",
        province="Misiones",
        city="Posadas",
        address="Ruta Excel",
    )
    carrier = Carrier.create(name="Transporte Excel")
    driver = Driver.create(name="Chofer Excel", carrier=carrier)
    truck = Truck.create(domain="XLSX123", carrier=carrier)
    nativa = Product.create(name="BOLSAS DE FECULA NATIVA", unit="UNIDAD")
    mezcla = Product.create(name="PACK FECULA X 1 KG", unit="UNIDAD")
    pallet_type = PalletType.create(type="Pallet Excel", measure="1x1", weight=0)
    order = LoadOrderService(current_user="excel").create_order(
        carrier=carrier,
        driver=driver,
        truck=truck,
        destinations=[
            {
                "client": client,
                "delivery_address": address,
                "products": [
                    {"product": nativa, "quantity": 550},
                    {"product": mezcla, "quantity": 20},
                ],
            }
        ],
        pallets=[
            {
                "sequence": 1,
                "pallet_type": pallet_type,
                "allocations": [
                    {
                        "client": client,
                        "delivery_address": address,
                        "product": nativa,
                        "quantity": 10,
                    },
                    {
                        "client": client,
                        "delivery_address": address,
                        "product": mezcla,
                        "quantity": 20,
                    },
                ],
            },
            *[
                {
                    "sequence": sequence,
                    "pallet_type": pallet_type,
                    "allocations": [
                        {
                            "client": client,
                            "delivery_address": address,
                            "product": nativa,
                            "quantity": 60,
                        }
                    ],
                }
                for sequence in range(5, 14)
            ],
        ],
    )

    target = LoadOrderPalletExcelExportService(current_user="excel").export(order, tmp_path)

    assert target.name == "orden_carga_1_armado_pallets.xlsx"
    workbook = load_workbook(target)
    sheet = workbook["Armado de pallets"]
    assert sheet["A1"].value == "2. DETALLE DEL PRODUCTO A DESPACHAR - ORDEN 0001"
    assert [sheet.cell(4, column).value for column in range(1, 6)] == [
        "Producto / detalle",
        "Cantidad total",
        "Pallet",
        "Lote",
        "Elab.",
    ]
    assert sheet["A6"].value == "BOLSAS DE FECULA NATIVA"
    assert sheet["B6"].value == "10 UNIDADES"
    assert sheet["C6"].value == "1"
    assert sheet["C7"].value is None
    assert "C6:C7" in {str(merged) for merged in sheet.merged_cells.ranges}
    assert sheet["A8"].value == "BOLSAS DE FECULA NATIVA"
    assert sheet["B8"].value == "60 UNIDADES"
    assert sheet["C8"].value == "9 pallets"
    assert AuditLog.select().where(AuditLog.action == "exportar_excel_pallets").count() == 1
